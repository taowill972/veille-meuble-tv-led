# -*- coding: utf-8 -*-
"""
build_excel_report.py
Génère le classeur Excel professionnel 'Classement_Meuble_TV_LED_AthisMons_2026.xlsx'
contenant 3 feuilles rigoureusement structurées et stylisées.
"""

import os
import sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.dirname(CURRENT_DIR)
DATA_DIR = os.path.join(BASE_DIR, "data")
OUTPUT_EXCEL = os.path.join(BASE_DIR, "Classement_Meuble_TV_LED_AthisMons_2026.xlsx")

def create_excel_report():
    print("=== Génération du Classeur Excel Professionnel Meuble TV LED ===")
    
    listings_file = os.path.join(DATA_DIR, "meuble_tv_led_listings.json")
    excluded_file = os.path.join(DATA_DIR, "meuble_tv_led_excluded.json")
    stats_file = os.path.join(DATA_DIR, "market_stats.json")

    with open(listings_file, "r", encoding="utf-8") as f:
        listings = json.load(f)
    with open(excluded_file, "r", encoding="utf-8") as f:
        excluded = json.load(f)
    stats = {}
    if os.path.exists(stats_file):
        with open(stats_file, "r", encoding="utf-8") as f:
            stats = json.load(f)

    wb = openpyxl.Workbook()
    # Retirer la feuille par défaut
    wb.remove(wb.active)

    # Styles
    font_title = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_cell = Font(name="Segoe UI", size=10)
    font_cell_bold = Font(name="Segoe UI", size=10, bold=True)
    font_link = Font(name="Segoe UI", size=10, underline="single", color="007AFF")
    font_note = Font(name="Segoe UI", size=9, italic=True, color="666666")

    fill_header = PatternFill(start_color="1A2B4C", end_color="1A2B4C", fill_type="solid") # Dark Navy
    fill_zebra = PatternFill(start_color="F7F9FC", end_color="F7F9FC", fill_type="solid")
    fill_accent = PatternFill(start_color="007AFF", end_color="007AFF", fill_type="solid")
    fill_top1 = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid") # Light emerald

    thin_border = Border(
        left=Side(style='thin', color="D0D7DE"),
        right=Side(style='thin', color="D0D7DE"),
        top=Side(style='thin', color="D0D7DE"),
        bottom=Side(style='thin', color="D0D7DE")
    )

    # ==========================================
    # FEUILLE 1 : TOP Classement LED
    # ==========================================
    ws1 = wb.create_sheet(title="TOP Classement LED")
    ws1.views.sheetView[0].showGridLines = True

    # Bannière Titre
    ws1.merge_cells("A1:M1")
    title_cell = ws1["A1"]
    title_cell.value = "🛋️ AUDIT & CLASSEMENT DÉCISIONNEL : MEUBLE TV À LED INTÉGRÉ (ATHIS-MONS 2026)"
    title_cell.font = font_title
    title_cell.fill = fill_accent
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 40

    # Sous-titre stats
    ws1.merge_cells("A2:M2")
    sub_cell = ws1["A2"]
    sub_cell.value = f"Rayon : ≤ 25 km Athis-Mons (91200) · Total analysé : {stats.get('total_analyzed', len(listings))} annonces · Médiane Marché : {stats.get('median_price', 75)} € · Mise à jour : {stats.get('updated_at', '17/09/2026')}"
    sub_cell.font = font_note
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 22

    headers = [
        "Rang", "Score /100", "Prix (€)", "Jauge Affaire", "Décote",
        "Titre de l'Annonce", "Ville & Distance", "Système LED & Contrôle",
        "Dimensions", "Finition & Style", "Rangements & Câbles", "Scénic III", "Lien Leboncoin"
    ]

    ws1.append([]) # Ligne 3 vide
    ws1.append(headers) # Ligne 4
    ws1.row_dimensions[4].height = 28

    for col_idx, h in enumerate(headers, 1):
        cell = ws1.cell(row=4, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for i, it in enumerate(listings, 1):
        row_num = 4 + i
        ws1.row_dimensions[row_num].height = 24

        link_url = it.get("url", "")
        rank_label = f"#{i} PÉPITE" if i == 1 else f"#{i}"
        
        row_values = [
            rank_label,
            f"{it.get('score')}/100",
            it.get("price"),
            it.get("bargain_gauge", "🟢 🟢 🟢 ⚪"),
            f"-{it.get('discount', 0)}%",
            it.get("title", ""),
            f"{it.get('city', '')} ({it.get('dist', '')} km)",
            it.get("led_type", ""),
            it.get("dimensions_str", ""),
            it.get("finish_name", ""),
            it.get("storage_desc", ""),
            it.get("scenic_status", ""),
            "Voir sur Leboncoin"
        ]

        ws1.append(row_values)

        for col_idx in range(1, len(row_values) + 1):
            cell = ws1.cell(row=row_num, column=col_idx)
            cell.font = font_cell
            cell.border = thin_border
            
            # Alignements
            if col_idx in [1, 2, 3, 4, 5, 12, 13]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Style TOP 1
            if i == 1:
                cell.fill = fill_top1
                cell.font = font_cell_bold
            elif i % 2 == 0:
                cell.fill = fill_zebra

            # Lien hypertexte
            if col_idx == 13 and link_url:
                cell.hyperlink = link_url
                cell.font = font_link

    # ==========================================
    # FEUILLE 2 : Grille de Notation (100 pts)
    # ==========================================
    ws2 = wb.create_sheet(title="Grille de Notation")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:D1")
    g_title = ws2["A1"]
    g_title.value = "BARÈME MULTI-CRITÈRES DE SCORING COGNITIF /100 (MEUBLE TV LED)"
    g_title.font = font_title
    g_title.fill = fill_header
    g_title.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 36

    g_headers = ["Pilier d'Évaluation", "Points Max", "Critères & Conditions Déterminantes", "Impact Ergonomique"]
    ws2.append([])
    ws2.append(g_headers)
    ws2.row_dimensions[3].height = 26

    for col_idx, h in enumerate(g_headers, 1):
        cell = ws2.cell(row=3, column=col_idx)
        cell.font = font_header
        cell.fill = fill_accent
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    criteria_data = [
        ("1. Système Éclairage LED & Contrôle", 25, "Ruban LED RGB avec télécommande (+25 pts) ; Rétroéclairage clips étagères verre (+20 pts) ; Ruban d'appoint mono (+18 pts).", "Ambiance lumineuse personnalisable, contrôle distant."),
        ("2. Rapport Qualité/Prix & Décote Argus", 25, "Prix ≤ 45€ (+25 pts) ; 46-75€ (+22 pts) ; 76-100€ (+18 pts) ; > 100€ (+10-14 pts). Décote vs MSRP.", "Rentabilité immédiate pour investissement locatif."),
        ("3. Format & Logistique Scénic III / TV 55\"", 20, "Longueur 120-140cm directe sans démontage (+15 pts) ; 141-160cm (+11 pts). Proximité ≤ 15 km (+5 pts).", "Chargement coffre autonome et support TV OLED 55 pouces."),
        ("4. Rangements Multimédia & Passe-Câbles", 15, "Niches ouvertes (Box/PS5) + Portes/Tiroirs fermés (+15 pts) ; Caissons fermés (+12 pts) ; Niches nues (+11 pts).", "Câbles invisibles et dissipation thermique box TV."),
        ("5. Finition & État Visuel HD", 15, "Blanc laqué brillant High Gloss (+15 pts) ; Bi-matière chêne & blanc (+14 pts) ; Noir laqué (+13 pts).", "Harmonie design scandinave / contemporain moderne.")
    ]

    for row_idx, r in enumerate(criteria_data, 4):
        ws2.append(list(r))
        ws2.row_dimensions[row_idx].height = 28
        for col_idx in range(1, 5):
            cell = ws2.cell(row=row_idx, column=col_idx)
            cell.font = font_cell
            cell.border = thin_border
            if col_idx == 2:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = font_cell_bold
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # ==========================================
    # FEUILLE 3 : Modèles Exclus & Justifications
    # ==========================================
    ws3 = wb.create_sheet(title="Modèles Exclus & Justifications")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:D1")
    e_title = ws3["A1"]
    e_title.value = "LOG DE SINCÉRITÉ TECHNIQUE : ANNONCES EXCLUES & JUSTIFICATIONS"
    e_title.font = font_title
    e_title.fill = fill_header
    e_title.alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 36

    e_headers = ["Identifiant", "Titre de l'Annonce", "Prix / Ville", "Motif Précis d'Élimination"]
    ws3.append([])
    ws3.append(e_headers)
    ws3.row_dimensions[3].height = 26

    for col_idx, h in enumerate(e_headers, 1):
        cell = ws3.cell(row=3, column=col_idx)
        cell.font = font_header
        cell.fill = PatternFill(start_color="B91C1C", end_color="B91C1C", fill_type="solid") # Dark Red
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for row_idx, it in enumerate(excluded, 4):
        ws3.append([
            it.get("id", ""),
            it.get("title", ""),
            f"{it.get('price', '')} € · {it.get('city', '')}",
            it.get("reason", "")
        ])
        ws3.row_dimensions[row_idx].height = 22
        for col_idx in range(1, 5):
            cell = ws3.cell(row=row_idx, column=col_idx)
            cell.font = font_cell
            cell.border = thin_border
            if col_idx in [1, 3]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Ajustement automatique de la largeur des colonnes pour toutes les feuilles
    for sheet in [ws1, ws2, ws3]:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if cell.row in [1, 2]: # Ignore merged titles
                    continue
                if len(val) > max_len:
                    max_len = len(val)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(OUTPUT_EXCEL)
    print(f"Classeur Excel sauvegardé avec succès : {OUTPUT_EXCEL}")

if __name__ == "__main__":
    create_excel_report()
